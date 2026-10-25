"""Document extraction. Locators are page/paragraph/row/line, never invented."""
import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, asdict
from html.parser import HTMLParser
from pathlib import Path

TEXT_EXT = {'.txt', '.md', '.rst', '.csv', '.tsv', '.json', '.jsonl', '.yaml', '.yml',
            '.html', '.htm', '.xml', '.sql', '.js', '.jsx', '.ts', '.tsx', '.vue', '.py', '.php'}
DOC_EXT = {'.pdf', '.docx', '.xlsx', '.xlsm', '.xls', '.pptx'}
UNSUPPORTED = {'.doc', '.ppt', '.rtf', '.odt', '.ods', '.zip', '.7z', '.rar', '.png', '.jpg', '.jpeg'}


@dataclass
class Unit:
    text: str
    locator: str


@dataclass
class Extracted:
    units: list
    warnings: list
    status: str = 'ok'

    def to_dict(self):
        return {'units': [asdict(u) for u in self.units], 'warnings': self.warnings, 'status': self.status}


def decode(data):
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        return data.decode('utf-16')
    for encoding in ('utf-8-sig', 'cp1251'):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode('utf-8', 'replace')


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ('style', 'script'):
            self.skip += 1
        if tag in ('p', 'div', 'li', 'tr', 'h1', 'h2', 'h3', 'br'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('style', 'script') and self.skip:
            self.skip -= 1
        if tag in ('p', 'div', 'li', 'tr'):
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def check_container(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries = z.infolist()
        if len(entries) > 6000 or sum(i.file_size for i in entries) > 100_000_000:
            raise ValueError('Архив документа слишком большой после распаковки')


def text_units(text, prefix='lines', max_chars=1800):
    units, buffer, start = [], [], 1
    for n, line in enumerate(text.splitlines(), 1):
        # Keep record boundaries; long minified lines are split with honest same-line locators.
        if not line.strip() or sum(len(x) for x in buffer) + len(line) > max_chars:
            if buffer:
                units.append(Unit('\n'.join(buffer), '%s:%d-%d' % (prefix, start, n - 1)))
                buffer = []
        if not line.strip():
            continue
        if not buffer:
            start = n
        if len(line) > max_chars:
            for offset in range(0, len(line), max_chars):
                units.append(Unit(line[offset:offset + max_chars], '%s:%d:chars:%d' % (prefix, n, offset)))
        else:
            buffer.append(line)
    if buffer:
        units.append(Unit('\n'.join(buffer), '%s:%d-%d' % (prefix, start, n)))
    return units


def structured_units(value, prefix='$', depth=0):
    # A record stays separate from its neighbours: no cross-customer/rate joins.
    if depth > 15:
        return [Unit(json.dumps(value, ensure_ascii=False)[:1800], prefix + ':depth_limit')]
    if isinstance(value, list):
        return [u for i, item in enumerate(value[:10000]) for u in structured_units(item, '%s[%d]' % (prefix, i), depth + 1)]
    if isinstance(value, dict):
        scalars = {k: v for k, v in value.items() if not isinstance(v, (dict, list))}
        included = set()
        for k, v in value.items():
            if isinstance(v, dict) and k.lower() in ('company', 'customer', 'client', 'supplier', 'contractor', 'заказчик', 'исполнитель') and isinstance(v.get('name'), str):
                scalars[k] = v['name']
                included.add(k)
            elif isinstance(v, list) and len(v) <= 100 and all(isinstance(x, (str, int, float, bool)) for x in v):
                scalars[k] = ', '.join(str(x) for x in v)
                included.add(k)
        result = [Unit('\n'.join('%s: %s' % (k, str(v).lower() if isinstance(v, bool) else v) for k, v in scalars.items()), prefix)] if scalars else []
        for k, v in value.items():
            if isinstance(v, (dict, list)) and k not in included:
                result.extend(structured_units(v, prefix + '.' + k, depth + 1))
        return result
    return [Unit(str(value), prefix)]


def code_units(text):
    """Read literal object boundaries without executing JS/TS, JSX or expressions."""
    stack, objects = [], []
    quote, comment, escaped = None, None, False
    i = 0
    while i < len(text):
        char, pair = text[i], text[i:i + 2]
        if comment == 'line':
            if char == '\n':
                comment = None
        elif comment == 'block':
            if pair == '*/':
                comment = None
                i += 1
        elif quote:
            if escaped:
                escaped = False
            elif char == '\\':
                escaped = True
            elif char == quote:
                quote = None
        elif pair in ('//', '/*'):
            comment = 'line' if pair == '//' else 'block'
            i += 1
        elif char in ('"', "'", '`'):
            quote = char
        elif char == '{':
            stack.append(i)
        elif char == '}' and stack:
            start = stack.pop()
            block = text[start:i + 1]
            if re.search(r'\bisOutstaff\s*:\s*(?:true|false)\b', block) and re.search(r'\bcompany\s*:', block):
                if not any(start <= a < b <= i + 1 for a, b in objects):
                    objects.append((start, i + 1))
        i += 1
    if not objects:
        return text_units(text)
    units, remainder = [], list(text)
    for start, end in sorted(objects):
        first = text.count('\n', 0, start) + 1
        last = text.count('\n', 0, end) + 1
        units.append(Unit(text[start:end], 'lines:%d-%d' % (first, last)))
        for i in range(start, end):
            if remainder[i] != '\n':
                remainder[i] = ' '
    units.extend(text_units(''.join(remainder)))
    return units


def extract(data, filename, max_units=10000, max_chars=2_000_000):
    ext = Path(filename).suffix.lower()
    warnings, units, status = [], [], 'ok'
    if data.startswith(b'version https://git-lfs.github.com/spec/'):
        return Extracted([], ['Git LFS pointer: binary not downloaded'], 'lfs_pointer')
    if ext in UNSUPPORTED:
        return Extracted([], ['Формат требует конвертации или OCR'], 'unsupported')
    if ext in {'.docx', '.xlsx', '.xlsm', '.pptx'}:
        check_container(data)
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(''):
            return Extracted([], ['PDF защищён паролем'], 'encrypted')
        if len(reader.pages) > 100:
            warnings.append('page_limit:100')
        for page, obj in enumerate(reader.pages[:100], 1):
            content = obj.extract_text() or ''
            if content.strip():
                units.extend(text_units(content, 'page:%d:lines' % page))
            else:
                warnings.append('page:%d:needs_ocr' % page)
        if not units:
            status = 'needs_ocr'
    elif ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        # Paragraph groups preserve the usual buyer/supplier introductory clause.
        units.extend(text_units('\n'.join(p.text for p in doc.paragraphs), 'paragraph'))
        for t, table in enumerate(doc.tables, 1):
            headers = [c.text for c in table.rows[0].cells] if table.rows else []
            for row, cells in enumerate(table.rows[1:], 2):
                units.append(Unit(' | '.join('%s: %s' % (h, c.text) for h, c in zip(headers, cells.cells)),
                                  'table:%d:row:%d' % (t, row)))
    elif ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book:
                if sheet.max_row and sheet.max_row > 10001:
                    warnings.append('sheet:%s:row_limit:10000' % sheet.title)
                if sheet.max_column and sheet.max_column > 60:
                    warnings.append('sheet:%s:column_limit:60' % sheet.title)
                rows = sheet.iter_rows(max_row=min(sheet.max_row or 10001, 10001), max_col=min(sheet.max_column or 60, 60), values_only=True)
                headers = next(rows, ())
                for rownum, row in enumerate(rows, 2):
                    pairs = ['%s: %s' % (headers[i] or 'column_%d' % (i + 1), x)
                             for i, x in enumerate(row) if x is not None]
                    if pairs:
                        units.append(Unit(' | '.join(pairs), 'sheet:%s:row:%d' % (sheet.title, rownum)))
            warnings.append('Excel: значения формул читаются только из сохранённого кэша; формулы не вычисляются')
        finally:
            book.close()
    elif ext == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows > 10001:
                    warnings.append('sheet:%s:row_limit' % sheet.name)
                headers = sheet.row_values(0) if sheet.nrows else []
                for row in range(1, min(sheet.nrows, 10001)):
                    units.append(Unit(' | '.join('%s: %s' % (h, x) for h, x in zip(headers[:60], sheet.row_values(row)[:60]) if x != ''),
                                      'sheet:%s:row:%d' % (sheet.name, row + 1)))
        finally:
            book.release_resources()
    elif ext == '.pptx':
        from xml.etree import ElementTree
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = sorted(n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n))
            for name in names[:100]:
                root = ElementTree.fromstring(archive.read(name))
                text = '\n'.join(node.text or '' for node in root.iter() if node.tag.endswith('}t'))
                units.extend(text_units(text, name))
            if len(names) > 100:
                warnings.append('slide_limit:100')
    else:
        text = decode(data)
        if '\x00' in text:
            return Extracted([], ['Бинарное содержимое в текстовом файле'], 'unsupported')
        if ext in ('.html', '.htm'):
            parser = HTMLText()
            parser.feed(text)
            normalized = '\n'.join(line.strip() for line in ''.join(parser.parts).splitlines() if line.strip())
            units = text_units(normalized, 'html_text_lines')
        elif ext in ('.csv', '.tsv'):
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel_tab if ext == '.tsv' else csv.excel
            rows = csv.reader(io.StringIO(text), dialect)
            headers = next(rows, [])
            for n, row in enumerate(rows, 2):
                if n > 10001:
                    warnings.append('row_limit:10000')
                    break
                units.append(Unit(' | '.join('%s: %s' % (headers[i] if i < len(headers) else 'column_%d' % i, val)
                                  for i, val in enumerate(row)), 'row:%d' % n))
        elif ext == '.json':
            try:
                units = structured_units(json.loads(text))
            except (ValueError, RecursionError):
                warnings.append('invalid_json:read_as_text')
                units = text_units(text)
        elif ext == '.jsonl':
            for n, line in enumerate(text.splitlines()[:10000], 1):
                try:
                    units.extend(structured_units(json.loads(line), 'line:%d' % n))
                except ValueError:
                    units.append(Unit(line, 'line:%d' % n))
        elif ext in ('.js', '.jsx', '.ts', '.tsx'):
            units = code_units(text)
        else:
            units = text_units(text)
    bounded, count = [], 0
    for unit in units:
        if len(bounded) >= max_units or count >= max_chars:
            warnings.append('text_limit')
            break
        value = unit.text[:min(5000, max_chars - count)]
        if len(value) < len(unit.text):
            warnings.append('unit_truncated:' + unit.locator)
        count += len(value)
        if value.strip():
            bounded.append(Unit(value, unit.locator))
    coverage_warnings = [w for w in warnings if w != 'Excel: значения формул читаются только из сохранённого кэша; формулы не вычисляются']
    if coverage_warnings and status == 'ok':
        status = 'partial'
    return Extracted(bounded, warnings, status)
