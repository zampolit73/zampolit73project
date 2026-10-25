import base64
import io
import json
import tempfile
import unittest
import urllib.error
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from irlix_intel.analyze import analyze, money_values, redact, safe_links, valid_inn10
from irlix_intel.extract import Unit, extract
from irlix_intel.exporter import export, safe_cell
from irlix_intel.github import GitHub, APIError, Pause, retry_delay, repo_name
from irlix_intel.pipeline import Pipeline, candidate
from irlix_intel.store import Store, key

SOURCE = {'repo': 'public/source', 'path': 'contracts/request.md', 'url': 'https://github.com/public/source/blob/' + 'a' * 40 + '/request.md'}
PAIR = 'IT-аутстаффинг Java Senior.\nЗаказчик: ООО «Ритейл»\nПодрядчик: ООО «Техно»\nСтавка: 3 500 руб./час, включая НДС.\nДоговор № AU-42 от 01.10.2026. Оплата в течение 30 дней.'


class AnalyzerTests(unittest.TestCase):
    def test_pair_is_claim_not_signed_contract(self):
        x = analyze(Unit(PAIR, 'lines:1-5'), SOURCE)
        self.assertEqual(x['customers'], ['ООО «Ритейл»'])
        self.assertEqual(x['suppliers'], ['ООО «Техно»'])
        self.assertTrue(x['explicit_pair'])
        self.assertEqual(x['grade'], 'B')
        self.assertEqual(x['review_status'], 'unreviewed')
        self.assertEqual(x['rates'][0]['min'], 3500)
        self.assertEqual(x['rates'][0]['unit'], 'hour')
        self.assertEqual(x['rates'][0]['VAT'], 'including_VAT')
        self.assertIn('AU-42', x['contract_numbers'])
        self.assertIn('01.10.2026', x['document_dates'])
        self.assertTrue(x['payment_terms'])

    def test_salary_never_client_rate(self):
        x = analyze(Unit('Вакансия Java. Зарплата: 250 000 руб./мес', 'row:2'), SOURCE)
        self.assertEqual(x['category'], 'hiring_signal')
        self.assertEqual(x['rates'][0]['kind'], 'salary')
        self.assertFalse(x['explicit_pair'])

    def test_unknown_rate_units_not_invented(self):
        x = money_values('Ставка: 3500')
        self.assertEqual(x[0]['currency'], 'unknown')
        self.assertEqual(x[0]['unit'], 'unknown')

    def test_rate_range(self):
        x = money_values('Ставка: 2 500–3 200 руб./час без НДС')[0]
        self.assertEqual((x['min'], x['max'], x['VAT']), (2500, 3200, 'without_VAT'))

    def test_total_is_not_hourly(self):
        x = money_values('Сумма договора: 10 000 000 руб.')[0]
        self.assertEqual(x['kind'], 'contract_total_or_budget')

    def test_anonymous_customer_not_guessed(self):
        x = analyze(Unit('Аутстафф Java.\nЗаказчик: не раскрыт, крупный банк РФ\nПодрядчик: ООО «Техно»', 'p1'), SOURCE)
        self.assertFalse(x['customers'])
        self.assertFalse(x['explicit_pair'])

    def test_blank_role_does_not_capture_next_line(self):
        x = analyze(Unit('Аутстафф Java\nЗаказчик:\nИсполнитель: ООО «Подрядчик»', 'p1'), SOURCE)
        self.assertEqual(x['customers'], [])
        self.assertEqual(x['suppliers'], ['ООО «Подрядчик»'])

    def test_negative_statement(self):
        x = analyze(Unit('Мы не используем аутстаффинг. Ищем Java разработчика.', 'p1'), SOURCE)
        self.assertEqual(x['category'], 'negative_outstaff_statement')
        self.assertLessEqual(x['score'], 15)

    def test_synthetic_not_lead(self):
        x = analyze(Unit('Синтетический учебный пример.\n' + PAIR, 'p1'), SOURCE)
        self.assertEqual((x['grade'], x['score']), ('X', 0))

    def test_logo_list_not_relationship(self):
        x = analyze(Unit('Аутстаффинг. Наши клиенты: Банк и Ритейл. Java.', 'p1'), SOURCE)
        self.assertEqual(x['grade'], 'C')
        self.assertFalse(x['explicit_pair'])

    def test_dedicated_team_not_proven_outstaff(self):
        x = analyze(Unit('Договор: выделенная команда Java.\nЗаказчик: ООО «Банк»\nПодрядчик: ООО «Разработка»', 'p1'), SOURCE)
        self.assertFalse(x['explicit_pair'])
        self.assertEqual(x['grade'], 'C')

    def test_multiple_parties_not_cross_joined(self):
        x = analyze(Unit(PAIR + '\nЗаказчик: ООО «Другой»', 'p1'), SOURCE)
        self.assertFalse(x['explicit_pair'])

    def test_aliases_dont_invent_roles(self):
        x = analyze(Unit('Аутстафф. Статья про Тинькофф и Java.', 'p1'), SOURCE, {'Т-Банк': ['Тинькофф']})
        self.assertIn('Т-Банк', x['entities'])
        self.assertEqual(x['customers'], [])

    def test_personal_contacts_and_tokens_redacted(self):
        value = redact('a@example.org +7 (999) 123-45-67 ghp_' + 'a' * 35)
        self.assertNotIn('example.org', value)
        self.assertNotIn('123-45', value)
        self.assertNotIn('ghp_', value)

    def test_external_links_drop_secret_query(self):
        self.assertEqual(safe_links('https://example.org/path?token=secret#frag'), ['https://example.org/path'])
        self.assertEqual(safe_links('https://user:password@example.org/path'), [])

    def test_inn_checksum(self):
        self.assertTrue(valid_inn10('7325170079'))
        self.assertFalse(valid_inn10('7325170070'))


class ExtractTests(unittest.TestCase):
    def test_csv_rows_do_not_merge(self):
        value = 'Заказчик;Запрос;Ставка\nООО «Первый»;Аутстафф Java;3500 руб./час\nООО «Второй»;Вакансия React;зарплата 250000 руб./мес\n'
        units = extract(value.encode(), 'requests.csv').units
        self.assertEqual(len(units), 2)
        a, b = [analyze(u, SOURCE) for u in units]
        self.assertNotIn('Второй', a['excerpt'])
        self.assertEqual(b['rates'][0]['kind'], 'salary')

    def test_json_records_do_not_merge(self):
        units = extract(json.dumps([{'client': 'A', 'service': 'outstaff Java'}, {'client': 'B', 'service': 'outstaff React'}]).encode(), 'x.json').units
        self.assertEqual(len(units), 2)
        self.assertEqual([analyze(x, SOURCE)['customers'] for x in units], [['A'], ['B']])

    def test_nested_json_company_and_stack_stay_together(self):
        raw = json.dumps([{'company': {'name': 'First'}, 'isOutstaff': True, 'techStack': ['React', 'TypeScript']}, {'company': {'name': 'Second'}, 'isOutstaff': False, 'techStack': ['Java']}]).encode()
        findings = [analyze(u, SOURCE) for u in extract(raw, 'data.json').units]
        findings = [f for f in findings if f]
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]['workplaces'], ['First'])
        self.assertEqual(findings[0]['category'], 'employee_outstaff_claim')
        self.assertEqual(findings[0]['stacks'], ['React'])

    def test_js_company_stays_with_correct_outstaff_flag(self):
        code = "const A = {company: {name: 'First'},\n\nisOutstaff: false,\n\ntechStack: ['React']};\nconst B = {company: {name: 'Second'},\n\nisOutstaff: true,\n\ntechStack: ['Java']};"
        findings = [analyze(u, SOURCE) for u in extract(code.encode(), 'experience.tsx').units]
        findings = [f for f in findings if f]
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]['workplaces'], ['Second'])
        self.assertEqual(findings[0]['stacks'], ['Java'])
        self.assertEqual(findings[0]['category'], 'employee_outstaff_claim')

    def test_type_definition_not_signal(self):
        self.assertIsNone(analyze(Unit('isOutstaff: boolean;\nconst sOutstaff = style({color: red});', 'p1'), SOURCE))

    def test_encoding_cp1251(self):
        self.assertIn('аутстафф', extract('аутстафф Java'.encode('cp1251'), 'x.txt').units[0].text)

    def test_lfs_not_treated_as_document(self):
        self.assertEqual(extract(b'version https://git-lfs.github.com/spec/v1\noid sha256:abc', 'x.pdf').status, 'lfs_pointer')

    def test_bad_zip_rejected(self):
        with self.assertRaises(Exception):
            extract(b'not a docx', 'x.docx')

    def test_limits_visible(self):
        result = extract(b'a\n\nb\n\nc', 'x.txt', max_units=1)
        self.assertEqual(result.status, 'partial')
        self.assertIn('text_limit', result.warnings)

    def test_docx_text_and_table(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph(PAIR)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Заказчик'
        table.cell(0, 1).text = 'Запрос'
        table.cell(1, 0).text = 'ООО «Клиент»'
        table.cell(1, 1).text = 'Аутстафф QA'
        stream = io.BytesIO()
        doc.save(stream)
        units = extract(stream.getvalue(), 'x.docx').units
        self.assertTrue(any('table:1:row:2' == u.locator for u in units))
        self.assertTrue(any(analyze(u, SOURCE)['explicit_pair'] for u in units if analyze(u, SOURCE)))

    def test_xlsx_rows(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(['Заказчик', 'Запрос', 'Ставка'])
        sheet.append(['Ритейл', 'аутстафф Java', '3000 руб./час'])
        sheet.append(['Банк', 'аутстафф QA', '2500 руб./час'])
        stream = io.BytesIO()
        book.save(stream)
        units = extract(stream.getvalue(), 'x.xlsx').units
        self.assertEqual(len(units), 2)
        self.assertNotIn('Банк', units[0].text)

    def test_blank_pdf_needs_ocr(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(100, 100)
        stream = io.BytesIO()
        writer.write(stream)
        self.assertEqual(extract(stream.getvalue(), 'scan.pdf').status, 'needs_ocr')


class FakeResponse:
    def __init__(self, value, headers=None):
        self.body = json.dumps(value).encode()
        self.headers = headers or {}
    def read(self, n):
        return self.body[:n]
    def __enter__(self):
        return self
    def __exit__(self, *args):
        pass


class APIClientTests(unittest.TestCase):
    def test_no_private_repo(self):
        client = GitHub(token='dummy')
        with patch.object(client, 'get', return_value={'private': True, 'visibility': 'private'}):
            with self.assertRaises(APIError):
                client.metadata('test/private')

    def test_budget_pause(self):
        client = GitHub(max_requests=0)
        with self.assertRaises(Pause):
            client.get('/rate_limit')

    def test_external_url_rejected(self):
        with self.assertRaises(ValueError):
            GitHub().get('https://example.org')

    def test_retry_headers(self):
        self.assertEqual(retry_delay({'Retry-After': '120'}), 120)
        self.assertEqual(retry_delay({'X-RateLimit-Remaining': '0', 'X-RateLimit-Reset': '200'}, now=100), 102)

    def test_429_long_pause_no_fast_retry(self):
        class Opener:
            calls = 0
            def open(self, *args, **kwargs):
                self.calls += 1
                raise urllib.error.HTTPError('https://api.github.com/rate_limit', 429, 'limited', {'Retry-After': '600'}, io.BytesIO(b'limited'))
        opener = Opener()
        with self.assertRaises(Pause):
            GitHub(opener=opener, wait_limit=1).get('/rate_limit')
        self.assertEqual(opener.calls, 1)

    def test_tree_truncation_fallback(self):
        client = GitHub()
        answers = [
            {'truncated': True, 'tree': []},
            {'tree': [{'type': 'tree', 'path': 'docs', 'sha': 'sub'}]},
            {'tree': [{'type': 'blob', 'path': 'contract.pdf', 'sha': 'blob', 'size': 7}]},
        ]
        with patch.object(client, 'get', side_effect=answers) as get:
            self.assertEqual(client.tree('a/b', 'root')[0]['path'], 'docs/contract.pdf')
            self.assertEqual(get.call_count, 3)

    def test_blob_limit(self):
        client = GitHub()
        with patch.object(client, 'get', return_value={'size': 5, 'encoding': 'base64', 'content': base64.b64encode(b'12345').decode()}):
            with self.assertRaises(APIError):
                client.blob('a/b', 'sha', 3)

    def test_repo_normalization(self):
        self.assertEqual(repo_name('https://github.com/OWNER/Repo.git'), 'owner/repo')
        with self.assertRaises(ValueError):
            repo_name('owner/repo/../../etc')

    def test_skip_secrets_and_vendor(self):
        self.assertIsNone(candidate('.env', {'.env'}))
        self.assertIsNone(candidate('node_modules/a/contract.md', set()))
        self.assertIsNotNone(candidate('docs/terms.pdf', set()))


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.store = Store(self.base / 'state.sqlite')
        self.args = SimpleNamespace(max_queries=5, search_pages=1, max_repos=5,
                                    max_files_per_repo=1, max_file_mb=5, refresh_hours=24)
    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def fake(self):
        class API:
            token = 'dummy'
            downloads = 0
            def metadata(self, repo):
                return {'private': False, 'default_branch': 'main'}
            def snapshot(self, repo, branch):
                return 'commit1', 'tree1'
            def tree(self, repo, tree):
                return [{'type': 'blob', 'path': n + '.md', 'sha': n, 'size': 200} for n in ('contract1', 'contract2')]
            def blob(self, repo, sha, limit):
                self.downloads += 1
                return PAIR.encode()
        return API()

    def test_resume_and_idempotence(self):
        api = self.fake()
        p = Pipeline(self.store, api, self.args)
        p.accept_repo('a/b')
        p.scan_repos()
        self.assertEqual(api.downloads, 1)
        self.assertEqual(self.store.rows('SELECT status FROM repos')[0]['status'], 'queued')
        p.scan_repos()
        self.assertEqual(api.downloads, 2)
        self.assertEqual(self.store.rows('SELECT status FROM repos')[0]['status'], 'done')
        p.scan_repos()
        self.assertEqual(api.downloads, 2)
        self.assertEqual(self.store.counts()['findings'], 2)

    def test_pause_does_not_finish_file(self):
        api = self.fake()
        p = Pipeline(self.store, api, self.args)
        p.accept_repo('a/b')
        with patch.object(api, 'blob', side_effect=Pause('limit')):
            with self.assertRaises(Pause):
                p.scan_repos()
        self.assertEqual(self.store.counts()['file_statuses'], {'pending': 2})
        p.scan_repos()
        self.assertEqual(self.store.counts()['findings'], 1)

    def test_tree_pause_does_not_commit_new_snapshot(self):
        api = self.fake()
        p = Pipeline(self.store, api, self.args)
        p.accept_repo('a/b')
        with patch.object(api, 'tree', side_effect=Pause('limit')):
            with self.assertRaises(Pause):
                p.scan_repos()
        self.assertIsNone(self.store.rows('SELECT commit_sha FROM repos')[0]['commit_sha'])
        p.scan_repos()
        self.assertEqual(self.store.counts()['findings'], 1)

    def test_search_page_committed_after_items(self):
        api = self.fake()
        response = {'total_count': 250, 'items': [{'private': False, 'full_name': 'x/r%d' % i} for i in range(100)]}
        api.search = lambda *args: response
        p = Pipeline(self.store, api, self.args)
        p.discover([{'kind': 'repositories', 'query': 'outstaff is:public'}])
        job = self.store.rows('SELECT * FROM search_jobs')[0]
        self.assertEqual((job['page'], job['done']), (2, 0))
        self.assertEqual(self.store.counts()['repos'], 100)

    def test_code_search_only_enqueues_candidates(self):
        api = self.fake()
        api.search = lambda *args: {'total_count': 2, 'items': [
            {'repository': {'full_name': 'a/b', 'private': False}, 'path': 'contract1.md'},
            {'repository': {'full_name': 'x/private', 'private': True}, 'path': 'secret.md'}]}
        p = Pipeline(self.store, api, self.args)
        p.discover([{'kind': 'code', 'query': 'аутстафф is:public'}])
        self.assertEqual(self.store.counts()['findings'], 0)
        self.assertEqual(self.store.counts()['repos'], 1)
        p.scan_repos()
        found = self.store.rows('SELECT * FROM findings')[0]
        self.assertIn('/blob/commit1/', found['url'])

    def test_public_rechecked_before_resume(self):
        api = self.fake()
        p = Pipeline(self.store, api, self.args)
        p.accept_repo('a/b')
        p.scan_repos()
        before = api.downloads
        p2 = Pipeline(self.store, api, self.args)
        with patch.object(api, 'metadata', side_effect=APIError(403, 'not public')):
            p2.scan_repos()
        self.assertEqual(before, api.downloads)
        self.assertEqual(self.store.rows('SELECT status FROM repos')[0]['status'], 'error')

    def test_incomplete_search_stays_pending(self):
        api = self.fake()
        api.search = lambda *args: {'total_count': 50, 'items': [], 'incomplete_results': True}
        Pipeline(self.store, api, self.args).discover([{'kind': 'code', 'query': 'outstaff'}])
        job = self.store.rows('SELECT * FROM search_jobs')[0]
        self.assertEqual((job['page'], job['done'], job['incomplete']), (1, 0, 1))

    def test_saturated_code_search_splits_without_gaps(self):
        api = self.fake()
        api.search = lambda *args: {'total_count': 2000, 'items': []}
        p = Pipeline(self.store, api, self.args)
        p.discover([{'kind': 'code', 'query': 'outstaff is:public'}])
        jobs = self.store.rows('SELECT * FROM search_jobs ORDER BY rowid')
        self.assertEqual(len(jobs), 3)
        self.assertEqual(jobs[0]['done'], 1)
        self.assertTrue(jobs[1]['query'].endswith('size:0..196607'))
        self.assertTrue(jobs[2]['query'].endswith('size:196608..393215'))
        self.assertEqual(jobs[1]['done'], 0)

    def test_contract_clause_roles(self):
        text = 'Договор аутстаффинга Java. ООО «Клиент», именуемое в дальнейшем «Заказчик», и ООО «Техно», именуемое в дальнейшем «Исполнитель».'
        found = analyze(Unit(text, 'p1'), SOURCE)
        self.assertEqual(found['customers'], ['ООО «Клиент»'])
        self.assertEqual(found['suppliers'], ['ООО «Техно»'])

    def test_export_formula_and_html_injection(self):
        fid = self.store.add_file('a/b', 'x.md', 'x', 'c', SOURCE['url'], 10)
        finding = analyze(Unit(PAIR + '\n<script>alert(1)</script>', 'p1'), SOURCE)
        finding['customers'] = ['=HYPERLINK("bad")']
        self.store.save_findings(fid, [finding], 'done', '', 'sha256', 1)
        out = self.base / 'report'
        export(self.store, out)
        page = (out / 'report.html').read_text()
        self.assertNotIn('<script>alert(1)</script>', page)
        self.assertIn('&lt;script&gt;alert(1)&lt;/script&gt;', page)
        self.assertEqual(safe_cell(' =2+2'), "' =2+2")
        from openpyxl import load_workbook
        book = load_workbook(out / 'findings.xlsx')
        self.assertFalse(any(c.data_type == 'f' for row in book['Находки'] for c in row))
        book.close()

    def test_review_survives_export_and_no_duplicate_findings(self):
        fid = self.store.add_file('a/b', 'x.md', 'x', 'c', SOURCE['url'], 10)
        f = analyze(Unit(PAIR, 'p1'), SOURCE)
        self.store.save_findings(fid, [f], 'done', '', 'sha', 1)
        self.store.execute('INSERT INTO reviews VALUES(?,?,?,?)', (key(fid, 'p1'), 'needs_followup', 'check', 'today'))
        self.store.save_findings(fid, [f], 'done', '', 'sha', 1)
        self.assertEqual(self.store.counts()['findings'], 1)
        self.assertEqual(len(self.store.rows('SELECT * FROM reviews')), 1)


if __name__ == '__main__':
    unittest.main()
